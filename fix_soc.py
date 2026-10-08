import docx

doc_path = r'C:\Users\LuanDias\PLANOS_LUAN_DADOS\PDF_AULAS\SOCIOLOGIA\EM\3_BIMESTRE\1_ANO_2_ANO_3_ANO\METODOLOGIA_CDP_ENSINO_MEDIO_MULTISSERIADO_J_3_B.docx'
doc = docx.Document(doc_path)

for p in doc.paragraphs:
    if p.text.startswith(('Para começar:', 'Foco no conteúdo:', 'Pause e responda:', 'Na prática:', 'Encerramento:')):
        if len(p.text) > 340:
            # We want to preserve the tag and bolding
            parts = p.text.split(': ', 1)
            tag = parts[0] + ': '
            content = parts[1]
            
            # split by '. '
            sentences = content.split('. ')
            
            # rebuild until length limit
            new_content = ""
            for s in sentences:
                addition = s + ('. ' if not s.endswith('.') else ' ')
                if len(tag) + len(new_content) + len(addition) > 340:
                    break
                new_content += addition
                
            new_content = new_content.strip()
            if not new_content:
                # if first sentence itself is too long
                new_content = content[:340 - len(tag) - 3] + '...'
            
            new_text = tag + new_content
            p.text = ''
            run = p.add_run(tag)
            run.bold = True
            p.add_run(new_content)

doc.save(doc_path)
