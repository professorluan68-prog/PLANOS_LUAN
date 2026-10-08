import docx

doc_path = r"C:\Users\LuanDias\PLANOS_LUAN_DADOS\PDF_AULAS\HISTÓRIACDP\EF\3_BIMESTRE\8_9_ANO_MULTISSERIADO\METODOLOGIA_HISTORIACDP_8_9_ANO_3_B.docx"
doc = docx.Document(doc_path)

replacements = {
    "Explicar o golpe civil-militar de 1964: deposição de João Goulart com apoio de setores empresariais e conservadores, sob pretexto de combater o comunismo. Apresentar os Atos Institucionais, destacando o AI-1 (1964): cassação de mandatos políticos, suspensão de direitos de opositores, fechamento de organizações e eleições indiretas.": "Explicar o golpe de 1964: deposição de Jango com apoio empresarial e conservador, sob pretexto de combater o comunismo. Apresentar os Atos Institucionais, com destaque ao AI-1: cassação de mandatos, suspensão de direitos de opositores, fechamento de organizações e eleições indiretas."
}

def normalize(text):
    return " ".join(text.split())

replacements_norm = {normalize(k): v for k, v in replacements.items()}

count = 0
for p in doc.paragraphs:
    if "Foco no conteúdo:" in p.text:
        parts = p.text.split(":", 1)
        if len(parts) == 2:
            prefix, content = parts[0], parts[1].strip()
            norm_content = normalize(content)
            for orig_key, new_val in replacements_norm.items():
                if orig_key == norm_content or orig_key in norm_content:
                    p.text = ""
                    run = p.add_run(prefix + ": ")
                    run.bold = True
                    p.add_run(new_val)
                    count += 1
                    break

print(f"Modificados {count} parágrafos extras.")
doc.save(doc_path)
