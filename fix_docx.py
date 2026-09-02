import docx
import re

doc_path = r'C:\Users\LuanDias\PLANOS_LUAN_DADOS\PDF_AULAS\LIDERANCA_E_ORATORIA\EM\3_BIMESTRE\2_ANO\METODOLOGIA_LIDERANCA_E_ORATORIA_2_ANO_3_B.docx'
doc = docx.Document(doc_path)

for p in doc.paragraphs:
    if p.text.startswith('Foco no conteúdo:'):
        match = re.search(r'Construir a reflexao sobre (.*?) por meio de exemplos', p.text)
        texto = match.group(1) if match else 'o tema principal'
        if len(texto) > 100:
            texto = texto[:97] + '...'
        new_text = f'Foco no conteúdo: Mediar leitura do material. Organizar no quadro ideias principais. Refletir sobre: {texto}. Usar exemplos cotidianos para relacionar sentir, pensar e agir.'
        if len(new_text) > 350:
             new_text = new_text[:347] + '...'
        p.text = ''
        run = p.add_run('Foco no conteúdo: ')
        run.bold = True
        p.add_run(new_text[18:])
            
    elif p.text.startswith('Na prática:'):
        new_text = 'Na prática: Realizar correção dialogada retomando material e dúvidas. Acompanhar atividade com registro individual. Garantir socialização opcional, evitando exposição.'
        if len(new_text) > 350:
             new_text = new_text[:347] + '...'
        p.text = ''
        run = p.add_run('Na prática: ')
        run.bold = True
        p.add_run(new_text[12:])
        
    elif p.text.startswith('Para começar:'):
        match = re.search(r'relacionada a \"(.*?)\",', p.text)
        tema = match.group(1) if match else 'o tema'
        if len(tema) > 100:
            tema = tema[:97] + '...'
        new_text = f'Para começar: Abrir a aula com situação acolhedora sobre \"{tema}\". Propor roda de conversa, respeitando ritmos e sem exigir exposição pessoal.'
        if len(new_text) > 350:
             new_text = new_text[:347] + '...'
        p.text = ''
        run = p.add_run('Para começar: ')
        run.bold = True
        p.add_run(new_text[14:])
        
    elif p.text.startswith('Encerramento:'):
        match = re.search(r'relacionado a \"(.*?)\",', p.text)
        tema = match.group(1) if match else 'o tema'
        if len(tema) > 100:
            tema = tema[:97] + '...'
        new_text = f'Encerramento: Concluir com observação para a semana sobre \"{tema}\", reforçando autonomia, respeito e cuidado nas relações.'
        if len(new_text) > 350:
             new_text = new_text[:347] + '...'
        p.text = ''
        run = p.add_run('Encerramento: ')
        run.bold = True
        p.add_run(new_text[14:])

doc.save(doc_path)
