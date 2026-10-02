"""
Módulo de renderização da revisão pedagógica e finalização de planos de aula.
"""
from __future__ import annotations

from collections import defaultdict
import hashlib
import os
from pathlib import Path
import re
import streamlit as st

from docx import Document

from core.normalizacao import normalizar as normalizar_texto_aux
from core.operacao import detectar_alteracoes_planos_revisados
from core.revisao_final import calcular_sha256, gravar_sidecar_json
from core.validador_plano import validar_aderencia_palavras_chave
from ui.relatorio_conferencia import _salvar_relatorios_conferencia
from ui.shared import _metodologia_app_para_blocos, _texto_metodologia_app


def renderizar_passo_revisao(
    *,
    professor: str,
    disciplina: str,
    disciplina_saida: str,
    turma: str,
    mes: str,
    bimestre: str,
    modo_ia: str,
    modo_upload_pdf: str,
    pasta_pdfs_auto: str,
    pdfs_selecionados_tela,
    modelo_bytes: bytes,
    escola: str,
    componente_curricular: str,
    semana: str,
    observacao: str,
    aulas_previstas_manual,
    fn_gerar_docx_final,
    fn_salvar_planos_na_pasta_finalizados,
    fn_salvar_planos_gerados_se_configurado,
    fn_registrar_mensagem_memoria_plano,
) -> tuple[list[dict], bool]:
    """
    Renderiza o Passo 2: Revisão pedagógica interativa, validação de palavras-chave,
    salvamento de referências e geração dos documentos Word finais.
    """
    turmas_processadas = st.session_state.get("turmas_processadas")
    if not turmas_processadas:
        return [], False

    avisos_processamento = st.session_state.get("avisos_processamento") or []
    for bloco in avisos_processamento:
        turma_aviso = str(bloco.get("turma") or "").strip()
        avisos_turma = [str(aviso).strip() for aviso in bloco.get("avisos", []) if str(aviso).strip()]
        if avisos_turma:
            prefixo = f"{turma_aviso}: " if turma_aviso else ""
            st.warning(prefixo + " | ".join(avisos_turma))

    st.markdown('<div class="section-title">✏️ Passo 2: Revisão</div>', unsafe_allow_html=True)
    st.markdown(
        '<div class="section-subtitle">Ajuste tema, aprendizagem, metodologia, acompanhamento e acessibilidade antes de montar o arquivo final.</div>',
        unsafe_allow_html=True,
    )

    total_turmas_revisao = len(turmas_processadas)
    total_aulas_revisao = sum(len(td.get("aulas", [])) for td in turmas_processadas)
    st.markdown(
        f"""
        <div class="review-shell">
            <div class="panel-title">Revisão pedagógica centralizada</div>
            <div class="panel-text">Você está revisando <strong>{total_aulas_revisao}</strong> aula(s) distribuídas em <strong>{total_turmas_revisao}</strong> turma(s). Abra apenas os blocos que quiser ajustar.</div>
            <div class="panel-pills">
                <span class="panel-pill">Tema</span>
                <span class="panel-pill">Metodologia</span>
                <span class="panel-pill">Acompanhamento</span>
                <span class="panel-pill">Acessibilidade</span>
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    turmas_revisadas = []
    rev_tok = st.session_state.get("revisao_token", 0)

    # Detecção de Frases Repetidas
    contagem_sentencas = defaultdict(list)
    for t_idx_dup, td_dup in enumerate(turmas_processadas):
        for a_idx_dup, aula_dup in enumerate(td_dup.get("aulas", [])):
            metodologia_dup = aula_dup.get("metodologia") or []
            textos_etapas = []
            for item in metodologia_dup:
                if isinstance(item, dict):
                    textos_etapas.append(item.get("texto", ""))
                else:
                    textos_etapas.append(str(item))
            texto_completo = " ".join(textos_etapas)
            sentencas = re.split(r"[.!?\n]", texto_completo)
            vistas_nesta_aula = set()
            for s in sentencas:
                s_limpa = re.sub(r"\s+", " ", s).strip()
                palavras = s_limpa.split()
                if len(palavras) > 8:
                    s_norm = normalizar_texto_aux(s_limpa)
                    if s_norm not in vistas_nesta_aula:
                        vistas_nesta_aula.add(s_norm)
                        contagem_sentencas[s_norm].append((t_idx_dup, a_idx_dup, s_limpa))

    duplicadas_por_aula = defaultdict(list)
    for frase_norm, ocorrencias in contagem_sentencas.items():
        if len(ocorrencias) > 2:
            for t_i, a_i, original_txt in ocorrencias:
                duplicadas_por_aula[(t_i, a_i)].append(original_txt)

    try:
        caminhos_relatorio = _salvar_relatorios_conferencia(
            turmas_processadas=turmas_processadas,
            duplicadas_por_aula=duplicadas_por_aula,
            professor=professor,
            disciplina=disciplina,
            turma=turma,
            mes=mes,
            bimestre=bimestre,
            modo_ia=modo_ia,
            modo_upload_pdf=modo_upload_pdf,
            pasta_pdfs_auto=pasta_pdfs_auto,
            pdfs_selecionados=pdfs_selecionados_tela,
        )
        if caminhos_relatorio:
            pasta_relatorio = Path(caminhos_relatorio[0]).parent
            st.info(
                "Relatórios de conferência salvos em: "
                f"{pasta_relatorio}. Esta pasta é só para análise e pode ser apagada depois sem afetar o sistema."
            )
    except Exception as err:
        st.warning(f"Não consegui salvar os relatórios de conferência automaticamente: {err}")

    for t_idx, td in enumerate(turmas_processadas):
        total_aulas_turma = len(td.get("aulas", []))
        st.markdown(f'<div class="review-class-title">{td["turma"]}</div>', unsafe_allow_html=True)
        st.markdown(
            f'<div class="review-class-meta">{total_aulas_turma} aula(s) prontas para conferência nesta turma.</div>',
            unsafe_allow_html=True,
        )
        aulas_edit = []
        for a_idx, aula in enumerate(td["aulas"]):
            score = aula.get("confidence_score", 100)
            if score >= 80:
                status_emoji = "🟢"
            elif score >= 60:
                status_emoji = "🟡"
            else:
                status_emoji = "🔴"

            with st.expander(f"{status_emoji} Aula {a_idx+1} - {aula.get('tema','')}", expanded=False):
                if score < 60:
                    st.error(f"🔴 **Qualidade Crítica ({score}%)**: Este plano possui baixíssima aderência ao PDF ou problemas pedagógicos graves.")
                elif score < 80:
                    st.warning(f"🟡 **Qualidade Aceitável ({score}%)**: O plano possui ressalvas ou desvios menores em relação ao PDF.")
                else:
                    st.success(f"🟢 **Alta Qualidade ({score}%)**: Plano totalmente aderente e validado.")

                avisos_val = aula.get("avisos_validacao") or []
                if avisos_val:
                    st.warning("**Alertas de Qualidade Pedagógica:**\n" + "\n".join([f"- {aviso}" for aviso in avisos_val]))

                frases_dupl = duplicadas_por_aula.get((t_idx, a_idx))
                if frases_dupl:
                    st.warning("**Aviso de Redundância (frases repetidas em mais de 2 aulas do lote):**\n" + "\n".join([f'- "{frase}"' for frase in frases_dupl]))

                t_val = st.session_state.get(f"tema_{rev_tok}_{t_idx}_{a_idx}")
                a_val = st.session_state.get(f"apr_{rev_tok}_{t_idx}_{a_idx}")
                acomp_val = st.session_state.get(f"acomp_{rev_tok}_{t_idx}_{a_idx}")
                aces_val = st.session_state.get(f"acess_{rev_tok}_{t_idx}_{a_idx}")
                m_val = st.session_state.get(f"met_{rev_tok}_{t_idx}_{a_idx}")

                if t_val is None:
                    t_val = aula.get("tema", "")
                if a_val is None:
                    a_val = aula.get("aprendizagem", "")
                if acomp_val is None:
                    acomp_val = "\n".join(aula.get("acompanhamento", []))
                if aces_val is None:
                    aces_val = "\n".join(aula.get("acessibilidade", []))
                if m_val is None:
                    m_val = _texto_metodologia_app(aula)

                palavras_chave_esperadas = aula.get("palavras_chave_esperadas") or []
                if palavras_chave_esperadas:
                    aula_temp = {
                        "metodologia": _metodologia_app_para_blocos(m_val),
                        "acompanhamento": [x.strip() for x in acomp_val.split("\n") if x.strip()],
                        "acessibilidade": [x.strip() for x in aces_val.split("\n") if x.strip()],
                    }

                    hash_content = f"{m_val}||{acomp_val}||{aces_val}||{','.join(palavras_chave_esperadas)}"
                    current_hash = hashlib.md5(hash_content.encode("utf-8")).hexdigest()

                    cache_key = f"pc_cache_{rev_tok}_{t_idx}_{a_idx}"
                    cache_data = st.session_state.get(cache_key)

                    if cache_data and cache_data.get("hash") == current_hash:
                        resultado_pc = cache_data["resultado"]
                    else:
                        resultado_pc = validar_aderencia_palavras_chave(aula_temp, palavras_chave_esperadas)
                        st.session_state[cache_key] = {
                            "hash": current_hash,
                            "resultado": resultado_pc,
                        }

                    cobertura_atual = resultado_pc["cobertura"]
                    valido_atual = resultado_pc["valido"]
                    palavras_ausentes_atuais = resultado_pc["palavras_ausentes"]

                    if valido_atual:
                        st.success(f"🎯 **Aderência de Palavras-Chave Validada ({cobertura_atual:.1f}%)**: Pelo menos 85% das palavras-chave obrigatórias estão presentes.")
                    else:
                        st.error(f"❌ **Plano Não Confiável - Aderência de Palavras-Chave Baixa ({cobertura_atual:.1f}%)**: O plano gerado não possui pelo menos 85% das palavras-chave obrigatórias.")
                        st.markdown("**Palavras-chave ausentes que devem ser incluídas no texto:**")
                        termos_ausentes_html = " ".join([f'<span style="background-color: #ffe6e6; color: #cc0000; padding: 2px 6px; border: 1px solid #ffcccc; border-radius: 4px; margin-right: 6px; font-family: monospace; font-size: 0.9em; display: inline-block; margin-bottom: 4px;">{palavra}</span>' for palavra in palavras_ausentes_atuais])
                        st.markdown(termos_ausentes_html, unsafe_allow_html=True)
                        st.caption("Dica: Edite os campos de Metodologia, Acompanhamento ou Acessibilidade abaixo e reinsira estes termos. O validador será atualizado instantaneamente.")
                else:
                    st.info("ℹ️ **Validação de Palavras-Chave**: Desativada no momento.")

                motivo_referencia_docx = str(
                    aula.get("motivo_referencia_docx") or ""
                ).strip()
                if motivo_referencia_docx:
                    st.warning("Referência DOCX desta aula: " + motivo_referencia_docx)

                col1, col2 = st.columns(2)
                with col1:
                    t = st.text_input("Tema", value=aula.get("tema", ""), key=f"tema_{rev_tok}_{t_idx}_{a_idx}")
                    a = st.text_area("Aprendizagem", value=aula.get("aprendizagem", ""), key=f"apr_{rev_tok}_{t_idx}_{a_idx}")
                with col2:
                    acomp = st.text_area("Acompanhamento", value="\n".join(aula.get("acompanhamento", [])), key=f"acomp_{rev_tok}_{t_idx}_{a_idx}")
                    aces = st.text_area("Acessibilidade", value="\n".join(aula.get("acessibilidade", [])), key=f"acess_{rev_tok}_{t_idx}_{a_idx}")
                m = st.text_area("Metodologia", value=_texto_metodologia_app(aula), height=150, key=f"met_{rev_tok}_{t_idx}_{a_idx}")

                # Relatório Técnico
                if st.checkbox("🛠️ Exibir Relatório Técnico da Geração", value=False, key=f"tech_rep_{rev_tok}_{t_idx}_{a_idx}"):
                    st.markdown(
                        f"""
                        | Parâmetro | Valor |
                        |---|---|
                        | **Provedor da IA** | {aula.get("ia_provedor") or "Sem IA"} |
                        | **Cache Reutilizado** | {"Sim" if aula.get("cache_reutilizado") else "Não"} |
                        | **Versão do Gerador** | {aula.get("versao_gerador") or "1.2.9"} |
                        | **Origem da Metodologia** | {aula.get("origem_metodologia") or "Desconhecida"} |
                        | **Score de Confiança** | {aula.get('confidence_score', 100)}% |
                        """
                    )

                    diag = aula.get("diagnostico_geracao") or {}
                    if diag:
                        st.markdown("#### Transformação da Metodologia (Pipeline)")
                        tabs = st.tabs(["1. Rascunho Local Heurístico", "2. Resposta IA Crua", "3. Higienização/Polimento", "4. Metodologia Final"])
                        with tabs[0]:
                            met_local = diag.get("metodologia_local") or []
                            if met_local:
                                st.write(_texto_metodologia_app({"metodologia": met_local}))
                            else:
                                st.info("Nenhuma etapa heurística local gerada.")
                        with tabs[1]:
                            met_ia = diag.get("metodologia_ia_crua") or []
                            if isinstance(met_ia, str):
                                st.text(met_ia)
                            elif met_ia:
                                st.write(_texto_metodologia_app({"metodologia": met_ia}))
                            else:
                                st.info("Sem resposta direta de IA (gerado localmente ou cached).")
                        with tabs[2]:
                            met_hig = diag.get("metodologia_higienizada") or []
                            if met_hig:
                                st.write(_texto_metodologia_app({"metodologia": met_hig}))
                            else:
                                st.info("Nenhum estágio higienizado intermediário.")
                        with tabs[3]:
                            met_fin = diag.get("metodologia_final") or []
                            if met_fin:
                                st.write(_texto_metodologia_app({"metodologia": met_fin}))
                            else:
                                st.info("Nenhuma metodologia final.")

                ae = aula.copy()
                ae.update({
                    "tema": t,
                    "aprendizagem": a,
                    "acompanhamento": [x.strip() for x in acomp.split("\n") if x.strip()],
                    "acessibilidade": [x.strip() for x in aces.split("\n") if x.strip()],
                    "metodologia": _metodologia_app_para_blocos(m),
                })
                palavras_chave_esperadas = ae.get("palavras_chave_esperadas") or []
                if palavras_chave_esperadas:
                    resultado_pc_final = validar_aderencia_palavras_chave(ae, palavras_chave_esperadas)
                    ae.update({
                        "valido_palavras_chave": resultado_pc_final["valido"],
                        "cobertura_palavras_chave": resultado_pc_final["cobertura"],
                        "palavras_chave_encontradas": resultado_pc_final["palavras_encontradas"],
                        "palavras_chave_ausentes": resultado_pc_final["palavras_ausentes"],
                    })
                aulas_edit.append(ae)
        turmas_revisadas.append({"turma": td["turma"], "aulas": aulas_edit})

    # Atualizar Arquivos de Referência DOCX
    referencias_para_atualizar = {}
    for tr in turmas_revisadas:
        for aula in tr["aulas"]:
            ref_path = aula.get("fonte_referencia_metodologia")
            if ref_path and os.path.exists(ref_path):
                referencias_para_atualizar.setdefault(ref_path, []).append(aula)

    if referencias_para_atualizar:
        st.markdown('<div class="section-card"></div><div class="section-title">💾 Atualizar Arquivos de Referência</div>', unsafe_allow_html=True)
        st.markdown('<div class="section-subtitle">Grave os ajustes e correções feitos nesta tela diretamente no arquivo DOCX de referência original.</div>', unsafe_allow_html=True)

        for ref_path, aulas_ref in referencias_para_atualizar.items():
            nome_ref_simpl = os.path.basename(ref_path)
            aulas_ref_unicas = {}
            for a in aulas_ref:
                num = a.get("numero_aula") or a.get("numero") or 0
                aulas_ref_unicas[num] = a

            btn_key = f"save_ref_{hashlib.md5(ref_path.encode('utf-8', errors='ignore')).hexdigest()[:8]}"
            confirmar_ref_key = f"confirm_ref_{hashlib.md5(ref_path.encode('utf-8', errors='ignore')).hexdigest()[:8]}"
            confirmar_ref = st.checkbox(
                f"Confirmo que desejo sobrescrever o DOCX de referência '{nome_ref_simpl}'.",
                key=confirmar_ref_key,
            )
            if st.button(
                f"Atualizar '{nome_ref_simpl}' com os ajustes desta tela",
                key=btn_key,
                type="secondary",
                disabled=not confirmar_ref,
            ):
                try:
                    doc = Document()
                    aulas_ordenadas = sorted(aulas_ref_unicas.values(), key=lambda x: int(x.get("numero_aula") or x.get("numero") or 0))
                    for aula in aulas_ordenadas:
                        num = aula.get("numero_aula") or aula.get("numero") or 0
                        tit = aula.get("tema") or ""
                        doc.add_paragraph(f"AULA {num} - {tit}")
                        doc.add_paragraph()
                        doc.add_paragraph("METODOLOGIA")
                        for etapa in (aula.get("metodologia") or []):
                            if isinstance(etapa, dict):
                                doc.add_paragraph(f"{etapa.get('titulo', '')}: {etapa.get('texto', '')}")
                            else:
                                doc.add_paragraph(str(etapa))
                        doc.add_paragraph()
                        doc.add_paragraph("ACOMPANHAMENTO DA APRENDIZAGEM")
                        for item in (aula.get("acompanhamento") or []):
                            item_limpo = str(item).replace("☑", "").strip()
                            if item_limpo:
                                doc.add_paragraph(f"☑ {item_limpo}")
                        doc.add_paragraph()
                        doc.add_paragraph("ACESSIBILIDADE")
                        for item in (aula.get("acessibilidade") or []):
                            item_limpo = str(item).replace("☑", "").strip()
                            if item_limpo:
                                doc.add_paragraph(f"☑ {item_limpo}")
                        doc.add_paragraph()

                    doc.save(ref_path)
                    st.success(f"✓ O arquivo '{nome_ref_simpl}' foi atualizado e agora contém as versões corrigidas dos planos!")
                except Exception as err:
                    st.error(f"Erro ao salvar arquivo de referência: {err}")

    # Atualizar o cache de metodologia base do PDF original (sidecar JSON)
    planos_com_pdf = []
    for tr in turmas_revisadas:
        for aula in tr["aulas"]:
            caminho_pdf_original = aula.get("caminho_pdf")
            if caminho_pdf_original and os.path.exists(caminho_pdf_original):
                planos_com_pdf.append(aula)

    if planos_com_pdf:
        st.markdown('<div class="section-card"></div><div class="section-title">🔄 Salvar no Cache de Metodologia Base (PDF)</div>', unsafe_allow_html=True)
        st.markdown('<div class="section-subtitle">Grave os ajustes feitos nesta tela no cache de metodologia base da aula original. As próximas gerações desta aula usarão esta versão corrigida automaticamente.</div>', unsafe_allow_html=True)

        confirmar_cache = st.checkbox(
            "Confirmo que desejo salvar as edições no cache de metodologia base permanente para estes PDFs.",
            key="confirmar_salvar_cache_base",
        )
        if st.button(
            "Salvar Alterações no Cache Base de Metodologias",
            key="btn_salvar_cache_base",
            type="secondary",
            disabled=not confirmar_cache,
        ):
            try:
                for aula in planos_com_pdf:
                    caminho_pdf_original = aula.get("caminho_pdf")
                    hash_pdf = aula.get("hash_pdf") or calcular_sha256(caminho_pdf_original)
                    gravar_sidecar_json(caminho_pdf_original, aula, hash_pdf)
                st.success("✓ O cache de metodologia base permanente foi atualizado com sucesso para todos os PDFs editados nesta tela!")
            except Exception as err:
                st.error(f"Erro ao salvar cache de metodologia base: {err}")

    st.markdown(
        """
        <div class="download-panel">
            <div class="panel-title">Última conferência antes do arquivo final</div>
            <div class="panel-text">Se estiver tudo certo na revisão, gere o documento final para liberar os botões de download logo abaixo.</div>
        </div>
        """,
        unsafe_allow_html=True,
    )
    st.checkbox(
        "Salvar este plano no histórico",
        key="salvar_historico_geracao",
        value=bool(st.session_state.get("salvar_historico_geracao", True)),
        help="Mantenha marcado para salvar este plano no histórico e habilitar a continuidade da sequência de PDFs na próxima geração.",
    )
    if st.button("GERAR DOCX", type="primary"):
        planos_gerados = []
        for tr in turmas_revisadas:
            planos_gerados.append(
                fn_gerar_docx_final(
                    modelo_bytes,
                    tr["aulas"],
                    escola,
                    professor,
                    disciplina,
                    componente_curricular,
                    tr["turma"],
                    mes,
                    bimestre,
                    semana,
                    observacao,
                    aulas_previstas_manual,
                )
            )
        st.session_state["planos_gerados"] = planos_gerados

        # Salva fisicamente na pasta do professor/disciplina
        fn_salvar_planos_na_pasta_finalizados(planos_gerados, disciplina_saida, professor, mes=mes)

        salvou_historico = fn_salvar_planos_gerados_se_configurado(
            planos_gerados,
            professor,
            disciplina_saida,
            bimestre,
            mes,
        )
        fn_registrar_mensagem_memoria_plano(salvou_historico)
        st.success("Planos gerados!")

    alteracoes_detectadas = detectar_alteracoes_planos_revisados(
        st.session_state.get("planos_gerados") or [],
        turmas_revisadas,
    )

    return turmas_revisadas, alteracoes_detectadas
