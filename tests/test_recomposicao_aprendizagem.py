"""
Testes automatizados para a formatação de Recomposição da Aprendizagem
em dias da semana configurados sem PDF.
"""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.shared import Pt, RGBColor
from docx_generator.preencher import (
    _titulo_aula,
    _preencher_celula_tema_material,
    _preencher_celula_recomposicao,
    _COR_VERMELHA,
)


def test_titulo_aula_bloco_sem_pdf():
    aula_sem_pdf = {
        "aula_vazia": True,
        "bloco_sem_pdf": True,
        "data": "10/11",
        "horario": "07h00 - 07h45",
    }
    titulo = _titulo_aula(aula_sem_pdf, numero=1)
    assert "RECOMPOSIÇÃO DA APRENDIZAGEM" in titulo
    assert "AULA:" in titulo
    assert "BIMESTRE" in titulo


def test_titulo_aula_vazia_antecipacao_permanece_vazia():
    aula_ant = {
        "aula_vazia": True,
        "bloco_sem_pdf": False,
        "data": "01/11",
        "horario": "07h00 - 07h45",
    }
    titulo = _titulo_aula(aula_ant, numero=1)
    assert titulo == ""


def test_formatacao_celula_recomposicao():
    doc = Document()
    tabela = doc.add_table(rows=1, cols=1)
    celula = tabela.cell(0, 0)

    _preencher_celula_recomposicao(celula)

    # Verifica os parágrafos gerados
    paras = celula.paragraphs
    assert len(paras) == 4

    # Parágrafo 1: RECOMPOSIÇÃO DA APRENDIZAGEM
    p1 = paras[0]
    assert p1.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert p1.text == "RECOMPOSIÇÃO DA APRENDIZAGEM"
    # O run com texto possui formatação vermelha e highlight amarelo
    runs_com_texto = [r for r in p1.runs if r.text.strip()]
    assert len(runs_com_texto) == 1
    assert runs_com_texto[0].bold is True
    assert runs_com_texto[0].font.color.rgb == _COR_VERMELHA
    assert runs_com_texto[0].font.highlight_color == WD_COLOR_INDEX.YELLOW
    assert runs_com_texto[0].font.size == Pt(9)

    # Parágrafo 2: Espaçamento em branco
    p2 = paras[1]
    assert p2.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert p2.text == ""

    # Parágrafo 3: AULA: 
    p3 = paras[2]
    assert p3.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert len(p3.runs) == 1
    assert "AULA:" in p3.runs[0].text
    assert p3.runs[0].bold is True
    assert p3.runs[0].font.color.rgb == _COR_VERMELHA
    assert p3.runs[0].font.size == Pt(9)

    # Parágrafo 4: BIMESTRE
    p4 = paras[3]
    assert p4.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert len(p4.runs) == 1
    assert "BIMESTRE" in p4.runs[0].text
    assert p4.runs[0].bold is True
    assert p4.runs[0].font.color.rgb == _COR_VERMELHA
    assert p4.runs[0].font.size == Pt(9)
