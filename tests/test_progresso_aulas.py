"""
Testes automatizados para a funcionalidade de progresso e continuidade de aulas.
"""
import pytest
from core.database import (
    salvar_progresso_aula,
    obter_progresso_aula,
    init_db,
    salvar_historico_plano,
)
from core.gestao_aulas import (
    obter_referencia_ultima_aula_ampla,
    obter_ultima_aula_gerada_sistema_impl,
)


def test_salvar_e_obter_progresso_aula(tmp_path):
    init_db()

    professor = "Teste Adriana Alda"
    disciplina = "Educação Financeira"
    turma = "6º ANO A"

    # Inicialmente pode ser None ou 0
    salvar_progresso_aula(professor, disciplina, turma, ultima_aula=14, ultimo_pdf="AULA 14.pdf", mes="Outubro")

    prog = obter_progresso_aula(professor, disciplina, turma)
    assert prog is not None
    assert prog["ultima_aula"] == 14
    assert prog["ultimo_pdf"] == "AULA 14.pdf"
    assert prog["mes_referencia"] == "Outubro"

    # Atualização para a próxima aula
    salvar_progresso_aula(professor, disciplina, turma, ultima_aula=22, ultimo_pdf="AULA 22.pdf", mes="Novembro")
    prog2 = obter_progresso_aula(professor, disciplina, turma)
    assert prog2 is not None
    assert prog2["ultima_aula"] == 22
    assert prog2["ultimo_pdf"] == "AULA 22.pdf"
    assert prog2["mes_referencia"] == "Novembro"


def test_obter_referencia_ultima_aula_ampla():
    init_db()

    professor = "Teste Professor Continuidade"
    disciplina = "Matemática"
    turma = "8º ANO B"

    # Salva no progresso
    salvar_progresso_aula(professor, disciplina, turma, ultima_aula=18)

    ref = obter_referencia_ultima_aula_ampla(professor, disciplina, turma, bimestre="4º Bimestre")
    assert ref["ultima_aula"] == 18
    assert ref["origem"] == "memoria_progresso"

    # Função de sistema impl deve retornar 18
    num = obter_ultima_aula_gerada_sistema_impl(professor, disciplina, turma, bimestre="4º Bimestre")
    assert num == 18


def test_obter_referencia_turmas_espelho_mesma_serie(tmp_path, monkeypatch):
    from docx import Document
    import config

    pasta_prof = tmp_path / "PROF_TESTE_ESPELHO" / "ARTE"
    pasta_prof.mkdir(parents=True, exist_ok=True)

    doc_a = Document()
    doc_a.add_paragraph("AULA 1 - Introdução")
    doc_a.add_paragraph("AULA 2 - Cores")
    doc_a.add_paragraph("AULA 3 - Formas")
    doc_a.save(pasta_prof / "Plano_1o_ANO_A_Arte.docx")

    doc_b = Document()
    doc_b.add_paragraph("AULA 1 - Introdução")
    doc_b.add_paragraph("AULA 2 - Cores")
    doc_b.add_paragraph("AULA 3 - Formas")
    doc_b.add_paragraph("AULA 4 - Movimento")
    doc_b.save(pasta_prof / "Plano_1o_ANO_B_Arte.docx")

    monkeypatch.setattr(config, "PLANOS_FEITOS_DIR", tmp_path)

    ref_a = obter_referencia_ultima_aula_ampla("PROF TESTE ESPELHO", "ARTE", "1º ANO A")
    ref_b = obter_referencia_ultima_aula_ampla("PROF TESTE ESPELHO", "ARTE", "1º ANO B")

    assert ref_a["ultima_aula"] == 4
    assert ref_b["ultima_aula"] == 4

