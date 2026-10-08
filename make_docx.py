import docx
from docx.shared import Pt
import re

# Cria o documento
doc = docx.Document()

# Ttulo principal
p = doc.add_paragraph()
p.add_run("HISTÓRIA – 1º, 2º E 3º ANOS DO ENSINO MÉDIO\n").bold = True
p.add_run("Metodologias, Acompanhamento da Aprendizagem e Acessibilidade\n").bold = True
p.add_run("3º Bimestre").bold = True

content = [
    {
        "aula": "1",
        "titulo": "Entre a fé e o controle: os mecanismos da Contrarreforma",
        "habilidade": "AE5 - Identificar as reformas religiosas e a Contrarreforma na Modernidade, relacionando-as à formação das monarquias nacionais, ao fortalecimento do absolutismo e ao mercantilismo.",
        "metodologia": {
            "Para começar:": "Iniciar com a palavra 'CENSURA' na lousa. Perguntar como as instituições controlam ideias perigosas à sua autoridade. Estimular a reflexão sobre o controle de informações, aproximando o tema da realidade adulta de vigilância e poder.",
            "Foco no conteúdo:": "Explicar a Contrarreforma Católica como resposta ao protestantismo. Destacar a Inquisição, o Index (livros proibidos) e a Companhia de Jesus. Usar a lousa para desenhar um esquema das estratégias da Igreja para manter seu poder e influência na Europa e colônias.",
            "Na prática:": "Leitura mediada de um pequeno texto impresso sobre um julgamento da Inquisição. No caderno, os estudantes respondem a duas questões sobre como o medo era usado para controle social e religioso. Correção e debate mediado pelo professor.",
            "Encerramento:": "Retomar as questões iniciais e pedir aos estudantes que sintetizem oralmente como o controle de ideias serve para manter estruturas de poder. O professor anota a conclusão principal na lousa."
        }
    },
    {
        "aula": "2",
        "titulo": "Poder divino e controle absoluto: a ascensão das monarquias europeias",
        "habilidade": "AE5 - Identificar as reformas religiosas e a Contrarreforma na Modernidade, relacionando-as à formação das monarquias nacionais, ao fortalecimento do absolutismo e ao mercantilismo.",
        "metodologia": {
            "Para começar:": "Escrever na lousa a famosa frase de Luís XIV: 'O Estado sou eu'. Questionar: o que acontece quando uma só pessoa concentra todo o poder, sem limites legais? Refletir sobre governos autoritários e a ausência de direitos.",
            "Foco no conteúdo:": "Abordar a formação dos Estados Nacionais e o Absolutismo monárquico. Explicar a Teoria do Direito Divino dos Reis, onde o monarca não devia explicações ao povo, apenas a Deus. Utilizar mapa mental na lousa conectando rei, exército e impostos.",
            "Na prática:": "Distribuição de um organograma impresso da sociedade absolutista. Os estudantes analisam e respondem no caderno por que o povo sustentava a nobreza e o clero sem ter direitos. Correção coletiva comentada.",
            "Encerramento:": "Debater a diferença entre ser 'súdito' (sem direitos, apenas deveres) e ser 'cidadão'. Fechar com a síntese de que a luta por leis escritas limitou o poder absoluto."
        }
    },
    {
        "aula": "3",
        "titulo": "Modernidade em movimento: entre permanências e rupturas",
        "habilidade": "AE5 - Identificar as reformas religiosas e a Contrarreforma na Modernidade, relacionando-as à formação das monarquias nacionais, ao fortalecimento do absolutismo e ao mercantilismo.",
        "metodologia": {
            "Para começar:": "Perguntar à turma: 'Quando dizemos que algo é moderno, o que queremos dizer?'. Discutir a transição do mundo rural (feudal) para o mundo urbano e comercial, relacionando com as transformações da vida adulta e do trabalho.",
            "Foco no conteúdo:": "Apresentar a Idade Moderna europeia como um período de transição: ascensão da burguesia, crescimento do comércio, Renascimento e novas formas de pensar, contrastando com as permanências da desigualdade social e exploração camponesa.",
            "Na prática:": "Atividade de associação no caderno. Os estudantes classificam elementos apresentados na lousa (comércio, trabalho rural, poder monárquico, ciência) como 'rupturas' ou 'permanências' em relação ao período feudal. Discussão e correção mediada.",
            "Encerramento:": "Construir uma frase-resumo na lousa: 'A modernidade trouxe inovações para alguns, mas manteve a exclusão de muitos'. Pedir que copiem o resumo no caderno."
        }
    },
    {
        "aula": "4",
        "titulo": "Da Metrópole à Colônia: a lógica mercantilista na expansão europeia",
        "habilidade": "AE5 - Identificar as reformas religiosas e a Contrarreforma na Modernidade, relacionando-as à formação das monarquias nacionais, ao fortalecimento do absolutismo e ao mercantilismo.",
        "metodologia": {
            "Para começar:": "Escrever na lousa: 'Lucro a qualquer custo'. Perguntar como a busca por riqueza imediata e exploração de recursos afeta as nações mais pobres. Relacionar as bases do sistema com as desigualdades atuais conhecidas pelos estudantes.",
            "Foco no conteúdo:": "Explicar o Mercantilismo como a política econômica dos reis absolutistas. Detalhar o colonialismo, o monopólio comercial, o pacto colonial e a exploração de metais preciosos e matérias-primas como formas de enriquecer a metrópole às custas da colônia.",
            "Na prática:": "Análise de um esquema impresso do 'Pacto Colonial'. Em seguida, os estudantes redigem um parágrafo no caderno explicando por que a colônia não conseguia se desenvolver economicamente. Leitura de algumas respostas e correção oral.",
            "Encerramento:": "Sintetizar com a turma como o sistema mercantilista deu origem às profundas desigualdades globais que perduram na atualidade. Registro coletivo final na lousa."
        }
    },
    {
        "aula": "5",
        "titulo": "Práticas e saberes ancestrais: a diversidade dos povos originários",
        "habilidade": "AE6 - Analisar a diversidade sociocultural e a organização territorial dos povos originários da América, relacionando-as à expansão marítima europeia e aos imaginários da colonização.",
        "metodologia": {
            "Para começar:": "Iniciar com a questão: 'A história do Brasil começou em 1500?'. Promover uma reflexão de que a terra já era habitada por milhões de pessoas, desmistificando a ideia de 'descobrimento' e reconhecendo as nações indígenas.",
            "Foco no conteúdo:": "Apresentar a diversidade de povos indígenas no Brasil (Tupi, Macro-Jê, entre outros) antes da colonização. Abordar sua organização social, cultura, línguas, relação com a natureza e ausência de noções europeias de propriedade privada e Estado.",
            "Na prática:": "Com base em um pequeno texto informativo impresso, os estudantes respondem a duas questões sobre o significado da terra para os povos originários em contraste com a visão mercantil europeia. Correção coletiva comentada.",
            "Encerramento:": "Reflexão mediada sobre como os saberes originários de respeito ao meio ambiente são cruciais no mundo de hoje. Registro de um apontamento final sobre a sabedoria ancestral no caderno."
        }
    },
    {
        "aula": "6",
        "titulo": "Ventos da mudança: navegações e conquistas nos séculos XV e XVI",
        "habilidade": "AE6 - Analisar a diversidade sociocultural e a organização territorial dos povos originários da América, relacionando-as à expansão marítima europeia e aos imaginários da colonização.",
        "metodologia": {
            "Para começar:": "Desenhar rapidamente uma caravela na lousa. Perguntar: 'Quais os riscos de enfrentar o desconhecido por interesses econômicos?'. Traçar um paralelo entre a ambição que impulsionou as navegações e os riscos de projetos humanos focados no lucro.",
            "Foco no conteúdo:": "Explicar as Grandes Navegações como uma grande empresa comercial para buscar especiarias e riquezas. Detalhar o protagonismo ibérico, as inovações (bússola, astrolábio) e a lógica de dominação e conquista do território 'descoberto'.",
            "Na prática:": "Leitura de um trecho impresso do diário de viagem de um navegador. Os alunos identificam no texto os interesses econômicos disfarçados de missão religiosa e anotam duas conclusões no caderno. Correção na lousa com o professor.",
            "Encerramento:": "Debater a conclusão: 'O que foi uma aventura comercial vitoriosa para a Europa representou o início de uma tragédia para a América e a África'. Síntese oral com a turma."
        }
    },
    {
        "aula": "7",
        "titulo": "Dois mundos em conflito: o encontro entre portugueses e povos originários",
        "habilidade": "AE6 - Analisar a diversidade sociocultural e a organização territorial dos povos originários da América, relacionando-as à expansão marítima europeia e aos imaginários da colonização.",
        "metodologia": {
            "Para começar:": "Escrever a palavra 'CHOQUE' na lousa. Perguntar o que acontece quando duas culturas completamente diferentes se encontram, sendo que uma possui armas de fogo e a intenção de lucrar.",
            "Foco no conteúdo:": "Abordar os primeiros contatos entre portugueses e tupis. Explicar o escambo (pau-brasil), o início da exploração e a transformação do contato inicial em violência, escravização e imposição cultural e religiosa (catequese).",
            "Na prática:": "Em folha impressa, os estudantes leem um trecho da Carta de Caminha e respondem a questões objetivas sobre como o europeu via os indígenas (curiosidade, ingenuidade e alvo de dominação). Correção conjunta com anotações na lousa.",
            "Encerramento:": "Retomar a ideia de choque cultural. Questionar: 'Foi um encontro ou um confronto?'. Registrar a conclusão da turma na lousa sobre a violência da colonização."
        }
    },
    {
        "aula": "8",
        "titulo": "Permanências de um imaginário: o Brasil indígena pelos olhos do colonizador",
        "habilidade": "AE6 - Analisar a diversidade sociocultural e a organização territorial dos povos originários da América, relacionando-as à expansão marítima europeia e aos imaginários da colonização.",
        "metodologia": {
            "Para começar:": "Colocar na lousa: 'O bom selvagem x O índio feroz'. Discutir como os preconceitos e estereótipos são criados para justificar a violência contra um povo, relacionando com estigmas presentes na sociedade atual.",
            "Foco no conteúdo:": "Explicar como os europeus construíram narrativas sobre os indígenas para justificar a conquista. Abordar o etnocentrismo (a Europa como modelo) e como essa visão inferiorizou as culturas locais para legitimar a tomada de terras e a escravização.",
            "Na prática:": "Atividade de verdadeiro ou falso no caderno sobre os mitos criados pelo colonizador em relação ao povo indígena. Em seguida, o professor promove a correção mediada, esclarecendo cada mito histórico.",
            "Encerramento:": "Debater como as visões preconceituosas do período colonial ainda afetam o respeito aos direitos indígenas no Brasil hoje. Sintetizar em uma frase de respeito e cidadania."
        }
    },
    {
        "aula": "9",
        "titulo": "Os donos da terra: território, poder e identidade nas civilizações inca, asteca e maia",
        "habilidade": "AE6 - Analisar a diversidade sociocultural e a organização territorial dos povos originários da América, relacionando-as à expansão marítima europeia e aos imaginários da colonização.",
        "metodologia": {
            "Para começar:": "Questionar a turma: 'Vocês sabiam que já existiam grandes impérios, com pirâmides e grandes capitais, na América antes da chegada de Colombo?'. Desconstruir o mito de que o continente era 'selvagem' e sem desenvolvimento civilizacional.",
            "Foco no conteúdo:": "Apresentar as civilizações pré-colombianas (Incas, Astecas e Maias). Destacar a complexidade política, a organização de grandes cidades (Tenochtitlán, Cusco), a agricultura avançada, sistemas de escrita, matemática e religião estatal.",
            "Na prática:": "Leitura compartilhada de um texto impresso sobre a infraestrutura e as ciências dos povos pré-colombianos. Os estudantes preenchem um quadro no caderno com os principais feitos (arquitetura, agricultura, astronomia). Correção na lousa.",
            "Encerramento:": "Relembrar a grandiosidade destas civilizações para combater a visão eurocêntrica da história. O professor anota uma frase-resumo sobre o desenvolvimento tecnológico e urbano da América originária."
        }
    },
    {
        "aula": "10",
        "titulo": "Materialidade e saberes incas: arquitetura, engenharia e cultura têxtil nos Andes",
        "habilidade": "AE6 - Analisar a diversidade sociocultural e a organização territorial dos povos originários da América, relacionando-as à expansão marítima europeia e aos imaginários da colonização.",
        "metodologia": {
            "Para começar:": "Perguntar: 'Como é possível construir estradas e pontes em montanhas muito altas (Cordilheira dos Andes) sem o maquinário moderno?'. Estimular a valorização da inteligência humana na superação de obstáculos naturais.",
            "Foco no conteúdo:": "Focar no Império Inca: a complexa rede de estradas (Caminho Inca), o sistema de comunicação, a agricultura em terraços nas montanhas e as técnicas de tecelagem que representavam status social e identidade cultural.",
            "Na prática:": "Atividade no caderno: relacionar os desafios geográficos das montanhas andinas com as soluções criadas pelos Incas (como os terraços agrícolas). Os alunos respondem duas questões práticas baseadas na explicação. Correção guiada.",
            "Encerramento:": "Destacar que o conhecimento de engenharia e adaptação ao clima hostil demonstra a imensa capacidade intelectual destes povos. Elaborar síntese oral junto com a turma."
        }
    },
    {
        "aula": "11",
        "titulo": "Conquista e Resistência: os impactos da colonização espanhola na América",
        "habilidade": "AE7 - Analisar as práticas colonizadoras espanhola e portuguesa na América, considerando as relações de poder, as estratégias de ocupação territorial e as dinâmicas entre europeus e indígenas.",
        "metodologia": {
            "Para começar:": "Iniciar com a pergunta: 'Como algumas centenas de europeus conseguiram dominar impérios de milhões de habitantes?'. Discutir fatores como a força bélica, as doenças, as alianças locais e a imposição brutal de poder.",
            "Foco no conteúdo:": "Explicar o processo de conquista espanhola. O uso de alianças com povos rivais, a disseminação de doenças europeias (varíola), a superioridade bélica e a posterior implementação da exploração do trabalho indígena (mita e encomienda).",
            "Na prática:": "Leitura de um texto impresso curto sobre a exploração nas minas de prata de Potosí. Os alunos respondem no caderno como o trabalho obrigatório imposto pelos colonizadores dizimou a população nativa. Correção e debate na lousa.",
            "Encerramento:": "Reflexão final sobre a exploração econômica e a resistência contínua dos povos nativos. Registrar que a riqueza europeia foi construída sobre o sacrifício forçado da América."
        }
    },
    {
        "aula": "12",
        "titulo": "Colonização portuguesa na América: poder, resistência e transformações sociais",
        "habilidade": "AE7 - Analisar as práticas colonizadoras espanhola e portuguesa na América, considerando as relações de poder, as estratégias de ocupação territorial e as dinâmicas entre europeus e indígenas.",
        "metodologia": {
            "Para começar:": "Escrever na lousa as palavras 'ENGENHO' e 'ESCRAVIDÃO'. Propor um debate sobre as bases que formaram a sociedade colonial brasileira, perguntando como a busca pelo lucro estruturou a desigualdade social no Brasil.",
            "Foco no conteúdo:": "Apresentar a colonização de exploração portuguesa: a monocultura da cana-de-açúcar, os latifúndios, o trabalho escravo (indígena e depois africano) e o modelo de exclusão social. Discutir também as formas de resistência negra e indígena.",
            "Na prática:": "Análise de um esquema na lousa sobre a estrutura do Engenho. Os alunos copiam o esquema e escrevem um parágrafo no caderno relacionando o uso do trabalho escravo à enorme concentração de riqueza. Correção coletiva comentada.",
            "Encerramento:": "Concluir a aula e o ciclo temático retomando que o passado colonial de exploração deixou marcas profundas na atual desigualdade e racismo estrutural da sociedade brasileira. Registro final com a turma."
        }
    }
]

acompanhamento = [
    "☑ Verificar se os estudantes compreenderam a relação entre os interesses de poder ou econômicos estudados e as ações dos grupos dominantes em cada contexto.",
    "☑ Observar a participação e a coerência das respostas orais e escritas, garantindo o entendimento individual das consequências dos processos históricos.",
    "☑ Acompanhar a leitura dos textos impressos e a resolução das questões, intervindo de forma clara para sanar dúvidas durante a atividade prática."
]

acessibilidade = [
    "☑ Realizar leitura mediada, expressiva e pausada de todos os textos impressos, garantindo a compreensão coletiva do vocabulário histórico e adulto.",
    "☑ Registrar palavras-chave, esquemas visuais e conceitos centrais na lousa com letra legível e tópicos bem espaçados para facilitar a cópia no caderno.",
    "☑ Permitir e encorajar que os estudantes utilizem a oralidade e relatos baseados em sua vivência prática para responder às questões e refletir."
]

for item in content:
    p = doc.add_paragraph()
    p.add_run(f"1º, 2º E 3º ANO – AULA {item['aula']} – {item['titulo']}\n").bold = True
    p.add_run(f"HABILIDADE: {item['habilidade']}")

    p = doc.add_paragraph()
    p.add_run("Metodologia").bold = True
    
    for chave, texto in item["metodologia"].items():
        p = doc.add_paragraph()
        p.add_run(f"{chave} ").bold = True
        p.add_run(texto)
        
    p = doc.add_paragraph()
    p.add_run("Acompanhamento da aprendizagem").bold = True
    for a in acompanhamento:
        doc.add_paragraph(a)
        
    p = doc.add_paragraph()
    p.add_run("Acessibilidade").bold = True
    for a in acessibilidade:
        doc.add_paragraph(a)

doc.save(r'C:\Users\LuanDias\PLANOS_LUAN_DADOS\PDF_AULAS\HISTÓRIACDP\EM\1_2_3_ANO\METODOLOGIA_HISTORIACDP_1_2_3_ANO_3_B.docx')
print("Documento DOCX criado com sucesso!")
